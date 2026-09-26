# -*- coding: utf-8 -*-
"""Generate SmartGift 3-sheet price workbook mock (Sheet1→2→3)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

OUT = Path(__file__).resolve().parents[1] / "public" / "ops" / "templates" / "price-sheet-3tab-template.xlsx"

THIN = Border(
    left=Side(style="thin", color="C4B8A8"),
    right=Side(style="thin", color="C4B8A8"),
    top=Side(style="thin", color="C4B8A8"),
    bottom=Side(style="thin", color="C4B8A8"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4D3A")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Tahoma", size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF8E7")
CALC_FILL = PatternFill("solid", fgColor="F3F1EC")
OK_FILL = PatternFill("solid", fgColor="E4F2E8")
BAD_FILL = PatternFill("solid", fgColor="FCE8E6")
TITLE = Font(name="Tahoma", size=14, bold=True, color="1F4D3A")
LABEL = Font(name="Tahoma", size=10)
NOTE = Font(name="Tahoma", size=9, color="666666", italic=True)


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build() -> Path:
    wb = Workbook()

    # ── Sheet1: สูตร / พารามิเตอร์ ─────────────────────────────────
    s1 = wb.active
    s1.title = "Sheet1_สูตร"
    s1["A1"] = "SmartGift · Sheet1 สูตรและพารามิเตอร์ (แก้ช่องเหลืองได้)"
    s1["A1"].font = TITLE
    s1.merge_cells("A1:F1")
    s1["A2"] = (
        "ชีตนี้เป็นสมองของสมุด — Sheet2 อ้าง FX / SOF / Markup จากที่นี่ "
        "แล้ว Sheet3 ดึงราคา Final จาก Sheet2 เท่านั้น"
    )
    s1["A2"].font = NOTE
    s1.merge_cells("A2:F2")

    s1["A4"] = "พารามิเตอร์"
    s1["B4"] = "ค่า"
    s1["C4"] = "หน่วย / หมายเหตุ"
    style_header(s1, 4, 3)

    params = [
        ("FX_CNY_THB", 5.0, "อัตราแลกเปลี่ยน (แก้ได้)"),
        ("โปรไฟล์", "standard", "standard หรือ corporate"),
        ("รวมค่าขนส่งจีน→ไทย", "YES", "YES / NO"),
        ("รวมแพ็กเกจ", "NO", "YES / NO · +80–150 บาทถ้า YES"),
        ("แพ็กเกจต่ำ", 80, "บาท/ชุด"),
        ("แพ็กเกจสูง", 150, "บาท/ชุด"),
        ("พื้นกำไร <500 ชุด", 5000, "บาท/ออเดอร์ (standard)"),
        ("พื้นกำไร 500+ ชุด", 3000, "บาท/ออเดอร์ (standard)"),
        ("พื้นกำไรองค์กร", 20000, "บาท · ใช้เมื่อโปรไฟล์=corporate"),
        ("Markup องค์กร", 1.47, "คงที่เมื่อโปรไฟล์=corporate"),
    ]
    for i, (name, value, note) in enumerate(params, start=5):
        s1.cell(i, 1, name).font = LABEL
        cell = s1.cell(i, 2, value)
        cell.fill = INPUT_FILL
        cell.border = THIN
        s1.cell(i, 3, note).font = NOTE

    # Named-ish anchors for formulas
    s1["E4"] = "สมการหลัก (อ่านอย่างเดียว)"
    s1["E4"].font = HEADER_FONT
    s1["E4"].fill = HEADER_FILL
    s1.merge_cells("E4:F4")
    s1["E5"] = "ลงเรือ"
    s1["F5"] = "โรงงานบาท + inland/ชิ้น + freight/ชิ้น"
    s1["E6"] = "ขาย"
    s1["F6"] = "ปัด(ลงเรือ × SOF × markup)"
    s1["E7"] = "โรงงานบาท"
    s1["F7"] = "โรงงานCNY × FX_CNY_THB"
    s1["E8"] = "หมายเหตุ"
    s1["F8"] = "บนเว็บใช้ computeUnitLanded จริง · ชีตนี้ประมาณด้วย inland/freight ที่ใส่ใน Sheet2"

    s1["A16"] = "ตาราง SOF (Small Order Factor)"
    s1["A16"].font = Font(name="Tahoma", bold=True, color="1F4D3A")
    s1["A17"] = "qty_max"
    s1["B17"] = "SOF"
    style_header(s1, 17, 2)
    sof_rows = [(20, 1.5), (50, 1.4), (100, 1.3), (300, 1.2), (499, 1.1), (999999, 1.0)]
    for i, (q, sof) in enumerate(sof_rows, start=18):
        s1.cell(i, 1, q).border = THIN
        s1.cell(i, 2, sof).border = THIN

    s1["D16"] = "ตาราง Markup ตามต้นทุนลงเรือ (บาท)"
    s1["D16"].font = Font(name="Tahoma", bold=True, color="1F4D3A")
    s1["D17"] = "landed_max"
    s1["E17"] = "markup"
    style_header(s1, 17, 2)
    # header styled only A17:B17 above — fix D17:E17
    for c in (4, 5):
        s1.cell(17, c).fill = HEADER_FILL
        s1.cell(17, c).font = HEADER_FONT
        s1.cell(17, c).border = THIN
    markup_rows = [(250, 3.0), (350, 2.73), (500, 2.62), (650, 2.45), (999999, 2.14)]
    for i, (cost, m) in enumerate(markup_rows, start=18):
        s1.cell(i, 4, cost).border = THIN
        s1.cell(i, 5, m).border = THIN

    s1["A25"] = "วิธีใช้: แก้ FX ที่ B5 → ไป Sheet2 แก้จำนวน (คอลัมน์ C) → ดูราคา Final ที่ Sheet3"
    s1["A25"].font = NOTE
    s1.merge_cells("A25:F25")
    autosize(s1, {1: 28, 2: 14, 3: 42, 4: 14, 5: 14, 6: 55})

    # ── Sheet2: สินค้า × ต้นทุน ───────────────────────────────────
    s2 = wb.create_sheet("Sheet2_ต้นทุน")
    s2["A1"] = "Sheet2 · รายการสินค้า × ต้นทุน (แก้ช่องเหลือง · คอลัมน์คำนวณอ้าง Sheet1)"
    s2["A1"].font = TITLE
    s2.merge_cells("A1:N1")
    s2["A2"] = (
        "คอลัมน์ inland/freight เป็นค่าประมาณต่อชิ้นเพื่อให้ลองสูตรใน Excel ได้ "
        "บนหน้าเว็บ /ops/price-sheet ระบบคำนวณขนส่งจากน้ำหนัก/CBM จริง"
    )
    s2["A2"].font = NOTE
    s2.merge_cells("A2:N2")

    headers = [
        "รหัส",
        "ชื่อสินค้า",
        "จำนวน (qty)",
        "โรงงาน CNY",
        "โรงงานบาท",
        "inland/ชิ้น",
        "freight/ชิ้น",
        "ลงเรือ",
        "SOF",
        "Markup",
        "ขาย/ชิ้น",
        "กำไรชุด",
        "พื้นกำไร",
        "ผ่านพื้น?",
    ]
    for c, h in enumerate(headers, start=1):
        s2.cell(4, c, h)
    style_header(s2, 4, len(headers))

    samples = [
        ("DEMO-BOX-01", "กล่องของขวัญพรีเมียม A", 50, 28, 8, 22),
        ("DEMO-TUM-02", "ชุดแก้วเก็บความเย็น", 100, 45, 10, 28),
        ("DEMO-HAM-03", "กระเช้าองค์กรพรีเมียม", 20, 80, 18, 45),
    ]

    for r, (code, name, qty, cny, inland, freight) in enumerate(samples, start=5):
        s2.cell(r, 1, code).border = THIN
        s2.cell(r, 2, name).border = THIN

        qty_cell = s2.cell(r, 3, qty)
        qty_cell.fill = INPUT_FILL
        qty_cell.border = THIN

        cny_cell = s2.cell(r, 4, cny)
        cny_cell.fill = INPUT_FILL
        cny_cell.border = THIN

        # E factory THB = D * Sheet1!B5
        e = s2.cell(r, 5, f"=D{r}*'Sheet1_สูตร'!$B$5")
        e.fill = CALC_FILL
        e.border = THIN
        e.number_format = "0.00"

        for col, val in ((6, inland), (7, freight)):
            cell = s2.cell(r, col, val)
            cell.fill = INPUT_FILL
            cell.border = THIN
            cell.number_format = "0.00"

        # H landed
        h = s2.cell(r, 8, f"=E{r}+F{r}+G{r}")
        h.fill = CALC_FILL
        h.border = THIN
        h.number_format = "0.00"

        # I SOF — same bands as SMALL_ORDER_FACTORS
        sof = s2.cell(
            r,
            9,
            (
                f"=IF(C{r}<=20,1.5,IF(C{r}<=50,1.4,IF(C{r}<=100,1.3,"
                f"IF(C{r}<=300,1.2,IF(C{r}<=499,1.1,1)))))"
            ),
        )
        sof.fill = CALC_FILL
        sof.border = THIN
        sof.number_format = "0.00"

        # J Markup — corporate override or same MARKUP_BANDS
        mk = s2.cell(
            r,
            10,
            (
                f'=IF(\'Sheet1_สูตร\'!$B$6="corporate",\'Sheet1_สูตร\'!$B$14,'
                f"IF(H{r}<=250,3,IF(H{r}<=350,2.73,IF(H{r}<=500,2.62,"
                f"IF(H{r}<=650,2.45,2.14)))))"
            ),
        )
        mk.fill = CALC_FILL
        mk.border = THIN
        mk.number_format = "0.00"

        # K sell = ROUND(H*I*J) + packaging if YES
        sell = s2.cell(
            r,
            11,
            (
                f'=ROUND(H{r}*I{r}*J{r},0)'
                f'+IF(\'Sheet1_สูตร\'!$B$8="YES",\'Sheet1_สูตร\'!$B$9,0)'
            ),
        )
        sell.fill = CALC_FILL
        sell.border = THIN
        sell.number_format = '#,##0'

        # L package profit
        profit = s2.cell(r, 12, f"=(K{r}-H{r})*C{r}")
        profit.fill = CALC_FILL
        profit.border = THIN
        profit.number_format = '#,##0'

        # M floor
        floor = s2.cell(
            r,
            13,
            (
                f'=IF(\'Sheet1_สูตร\'!$B$6="corporate",\'Sheet1_สูตร\'!$B$13,'
                f'IF(C{r}<500,\'Sheet1_สูตร\'!$B$11,\'Sheet1_สูตร\'!$B$12))'
            ),
        )
        floor.fill = CALC_FILL
        floor.border = THIN
        floor.number_format = '#,##0'

        # N pass?
        passed = s2.cell(r, 14, f'=IF(L{r}>=M{r},"PASS","BELOW")')
        passed.fill = CALC_FILL
        passed.border = THIN

    s2.conditional_formatting.add(
        "N5:N7",
        FormulaRule(formula=['$N5="PASS"'], fill=OK_FILL),
    )
    s2.conditional_formatting.add(
        "N5:N7",
        FormulaRule(formula=['$N5="BELOW"'], fill=BAD_FILL),
    )

    s2["A9"] = "ลองเปลี่ยนจำนวนในคอลัมน์ C แล้วดู SOF / ขาย / Sheet3 อัปเดตตาม"
    s2["A9"].font = NOTE
    autosize(
        s2,
        {
            1: 14,
            2: 28,
            3: 12,
            4: 12,
            5: 12,
            6: 12,
            7: 12,
            8: 12,
            9: 8,
            10: 10,
            11: 12,
            12: 12,
            13: 12,
            14: 10,
        },
    )

    # ── Sheet3: Final ────────────────────────────────────────────
    s3 = wb.create_sheet("Sheet3_Final")
    s3["A1"] = "Sheet3 · ราคา Final (พรีวิวลูกค้า — อ้างอิง Sheet2 อย่างเดียว)"
    s3["A1"].font = TITLE
    s3.merge_cells("A1:G1")
    s3["A2"] = "ห้ามแก้ราคา Final โดยตรง · แก้ qty/ต้นทุนที่ Sheet2 หรือ FX ที่ Sheet1"
    s3["A2"].font = NOTE
    s3.merge_cells("A2:G2")

    final_headers = [
        "ลำดับ",
        "รหัส",
        "ชื่อสินค้า",
        "จำนวน",
        "ราคา/ชุด (Final)",
        "รวมทั้งชุด",
        "สถานะพื้นกำไร",
    ]
    for c, h in enumerate(final_headers, start=1):
        s3.cell(4, c, h)
    style_header(s3, 4, len(final_headers))

    for i, r in enumerate(range(5, 8), start=1):
        s3.cell(r, 1, i).border = THIN
        s3.cell(r, 2, f"=Sheet2_ต้นทุน!A{r}").border = THIN
        s3.cell(r, 3, f"=Sheet2_ต้นทุน!B{r}").border = THIN
        s3.cell(r, 4, f"=Sheet2_ต้นทุน!C{r}").border = THIN
        price = s3.cell(r, 5, f"=Sheet2_ต้นทุน!K{r}")
        price.border = THIN
        price.number_format = '#,##0" บาท"'
        price.fill = OK_FILL
        total = s3.cell(r, 6, f"=E{r}*D{r}")
        total.border = THIN
        total.number_format = '#,##0" บาท"'
        s3.cell(r, 7, f"=Sheet2_ต้นทุน!N{r}").border = THIN

    s3["A9"] = "รวม Final ทั้งใบ"
    s3["A9"].font = Font(name="Tahoma", bold=True)
    s3["F9"] = "=SUM(F5:F7)"
    s3["F9"].font = Font(name="Tahoma", bold=True, color="1F4D3A", size=12)
    s3["F9"].number_format = '#,##0" บาท"'

    s3["A11"] = (
        "พรีวิวบนเว็บ: /ops/price-sheet (แก้ค่าได้ · คำนวณด้วย computeUnitLanded ชุดเดียวกับแคตตาล็อก)"
    )
    s3["A11"].font = NOTE
    s3.merge_cells("A11:G11")

    autosize(s3, {1: 8, 2: 14, 3: 28, 4: 10, 5: 18, 6: 16, 7: 14})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
